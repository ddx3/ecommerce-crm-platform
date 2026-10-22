# -*- coding: utf-8 -*-
"""A batch of instructions in one .docx.

Why Word, not PDF: the owner enables duplex printing precisely in Word.

And this gives what Chromium can't do at all. Each instruction is a separate SECTION starting "from an
odd page" (WD_SECTION.ODD_PAGE). Word counts the pages itself and itself
inserts a blank back side where an instruction ended on an odd number of pages — so each starts on
the front side of a fresh sheet. In PDF this had to be done by hand: render, count pages and stitch
in blanks (see instruction_print._merge). Chromium ignores `break-before: recto`, Word doesn't.

The order and number of sheets come from `instruction_print.collect_sheets` — the same function that
counted them for PDF: one sheet per (order, bundle) pair, identical bundles consecutively, in the
order of the "Packing" list.

The instruction HTML is parsed manually (h1/h3/p/ul/ol/li/strong/u/br) — exactly the set of tags the
CRM editor produces. We don't use a third-party conversion library: it drags in a dependency for a
dozen tags and still loses styles.
"""
from __future__ import annotations

import logging
import re
from html.parser import HTMLParser
from io import BytesIO

log = logging.getLogger(__name__)

_BLOCK = {'p', 'h1', 'h2', 'h3', 'h4', 'li', 'div'}
_SKIP = {'script', 'style', 'head', 'title'}

# Line-break marker inside a paragraph. In Word it is a separate <w:br/> element, not a character in
# the text: Word collapses '\n' inside a run into a space, and two lines the manager separated with
# Shift+Enter would merge into one.
_BR = '\x00br\x00'

# We remove decorative emoji from the text. The ranges cover pictograms only: the en dash, em dash
# and ellipsis (– — …) are NOT emoji and must not be thrown out — they hold together "8–10 L"
# and "+18…+25°C".
#
# 🐝 (U+1F41D) is an exception: in this catalog the bee means 'safe for bees' (Mavrik, Charivnyk),
# i.e. it carries meaning, not decoration. Removing it would lose the safety warning.
_BEE = '\U0001F41D'
_EMOJI_RE = re.compile(
    '(?![' + _BEE + '])['
    '\U0001F300-\U0001F5FF'   # symbols and pictograms
    '\U0001F600-\U0001F64F'   # smileys
    '\U0001F680-\U0001F6FF'   # transport and map symbols
    '\U0001F900-\U0001FAFF'   # supplemental
    '☀-⛿'           # miscellaneous symbols (☀ ⚠ …)
    '✀-➿'           # dingbats (✅ ✂ …)
    '⬀-⯿'           # arrow pictograms
    # Standalone signs from «General Punctuation» that are used as emoji.
    # This is where «‼» lands, which is where the problems started.
    '‼⁉™ℹ↔-↪⌚⌛⌨⏏'
    '⏩-⏺Ⓜ▪▫▶◀◻-◾'
    ']️?'
)
# After cleaning, orphans remain: the emoji presentation selector, the joiner and the zero-width
# character. We remove the zero-width one on its own too — it comes in from catalog names ("Koragen
# 🐝") and invisibly breaks text search.
_EMOJI_TAIL_RE = re.compile('[︎️‍​]')


def strip_emoji(text: str, *, trim: bool = True) -> str:
    """Remove decorative emoji, leaving 🐝 and punctuation.

    The en dash, em dash and ellipsis are NOT emoji: "8–10 L" and
    «+18…+25°C» rely on them, so the ranges are chosen not to touch them.

    `trim=False` — for a separate piece of a paragraph: there the trailing space is significant, it
    separates a bold word from the rest of the sentence, and it can't be trimmed.
    """
    out = _EMOJI_RE.sub('', text or '')
    out = _EMOJI_TAIL_RE.sub('', out)
    out = re.sub(r' {2,}', ' ', out)
    return out.strip() if trim else out


class _Node:
    """One paragraph of the future document: type + text pieces with styles."""

    __slots__ = ('kind', 'runs', 'level')

    def __init__(self, kind, level=0):
        self.kind = kind          # 'p' | 'h1' | 'h2' | 'bullet' | 'number'
        self.runs = []            # [(text, bold, underline)]
        self.level = level

    @property
    def text(self):
        return ''.join(t for t, _, _ in self.runs)

    @property
    def visible(self):
        """Text without break markers — so a 'paragraph made of only <br>' counts as empty."""
        return ''.join(t for t, _, _ in self.runs if t != _BR)


class _Reader(HTMLParser):
    """HTML -> flat list of paragraphs."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.nodes = []
        self._cur = None
        self._bold = 0
        self._under = 0
        self._skip = 0
        self._lists = []          # stack: 'ul' | 'ol'
        self._counter = 0         # item number of the current <ol>

    def _open(self, kind):
        self._flush()
        self._cur = _Node(kind)

    def _flush(self):
        if self._cur is None:
            return
        if self._cur.visible.strip():
            self.nodes.append(self._cur)
        elif self._cur.runs:
            # <p><br></p> — an empty line the manager uses to separate blocks from each other. It can't be thrown away: that is exactly why the whole sheet merged into solid text. We keep it as an empty paragraph.
            self._cur.kind = 'spacer'
            self._cur.runs = []
            self.nodes.append(self._cur)
        self._cur = None

    def handle_starttag(self, tag, attrs):
        if tag in _SKIP:
            self._skip += 1
            return
        if self._skip:
            return
        if tag in ('b', 'strong'):
            self._bold += 1
        elif tag in ('u', 'ins'):
            self._under += 1
        elif tag in ('ul', 'ol'):
            self._flush()
            self._lists.append(tag)
            self._counter = 0        # each <ol> is numbered from one
        elif tag == 'li':
            ordered = bool(self._lists and self._lists[-1] == 'ol')
            self._open('number' if ordered else 'bullet')
            if ordered:
                self._counter += 1
                self._cur.level = self._counter
        elif tag in ('h1',):
            self._open('h1')
        elif tag in ('h2', 'h3', 'h4'):
            self._open('h2')
        elif tag in ('p', 'div'):
            self._open('p')
        elif tag == 'br':
            if self._cur is None:
                self._cur = _Node('p')
            self._cur.runs.append((_BR, False, False))

    def handle_endtag(self, tag):
        if tag in _SKIP:
            self._skip = max(0, self._skip - 1)
            return
        if self._skip:
            return
        if tag in ('b', 'strong'):
            self._bold = max(0, self._bold - 1)
        elif tag in ('u', 'ins'):
            self._under = max(0, self._under - 1)
        elif tag in ('ul', 'ol'):
            self._flush()
            if self._lists:
                self._lists.pop()
        elif tag in _BLOCK:
            self._flush()

    def handle_data(self, data):
        if self._skip or not data:
            return
        if self._cur is None:
            # Text outside a block (a bare string at the root) is a paragraph too.
            self._cur = _Node('p')
        # Newlines inside HTML are not line breaks in Word: it has its own line length.
        text = data.replace('\r', ' ').replace('\n', ' ')
        while '  ' in text:
            text = text.replace('  ', ' ')
        if not text.strip() and not self._cur.runs:
            return
        self._cur.runs.append((text, self._bold > 0, self._under > 0))

    def close(self):
        super().close()
        self._flush()


def _promote_headings(nodes):
    """A paragraph that is ENTIRELY bold is a subheading, not a text paragraph.

    In real instructions the sections are formatted exactly like this: "TREATMENT No. 1 — SEPTEMBER",
    "Disease protection:", "PREPARATION SEQUENCE" — a separate bold line. There is no formal <h3> there, so
    without this step such a line gets the indent of an ordinary paragraph and sticks to the
    previous text — the sheet reads as a solid wall of text where you can't see where one treatment
    ended.

    We don't touch long lines: sometimes a whole warning phrase is highlighted in bold, and turning
    it into a heading would be wrong.
    """
    for n in nodes:
        if n.kind != 'p':
            continue
        runs = [(t, b) for t, b, _ in n.runs if t != _BR and t.strip()]
        if not runs or len(n.visible.strip()) > 90:
            continue
        if all(bold for _, bold in runs):
            n.kind = 'h2'
    return nodes


def parse_html(html: str) -> list:
    r = _Reader()
    try:
        r.feed(html or '')
        r.close()
    except Exception:
        log.exception('instruction html parse failed')
    return _promote_headings(r.nodes)


# ---------------------------------------------------------------- document

def _setup(doc):
    """A4 and margins for the same density the PDF had."""
    from docx.shared import Pt, Cm

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.1)
        section.bottom_margin = Cm(1.1)
        section.left_margin = Cm(1.2)
        section.right_margin = Cm(1.2)

    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(9.5)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15


# Paragraph spacing. The key thing here is `before` on headings: it is what sets a section off from
# the preceding text. Without it the sheet reads as a solid wall of text, and the person working at
# the plant can't tell where one treatment ended and the next began.
_STYLE = {
    'h1': {'size': 13, 'bold': True, 'before': 0, 'after': 7},
    'h2': {'size': 10.5, 'bold': True, 'before': 10, 'after': 3},
    'p': {'size': None, 'bold': False, 'before': 0, 'after': 4},
    'bullet': {'size': None, 'bold': False, 'before': 0, 'after': 2},
    'number': {'size': None, 'bold': False, 'before': 0, 'after': 2},
    'spacer': {'size': None, 'bold': False, 'before': 0, 'after': 0},
}


def _write_node(doc, node):
    from docx.shared import Pt

    if node.kind == 'spacer':
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        return p

    if node.kind == 'number':
        # We put the number ourselves, not via the 'List Number' style: Word numbers continuously
        # across the whole document, and the second "Mixing order" started at
        # 6. Our own counter resets at every <ol> — both within a sheet and between sheets.
        from docx.shared import Cm
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        node.runs.insert(0, (f'{node.level or 1}. ', False, False))
    else:
        style_name = 'List Bullet' if node.kind == 'bullet' else None
        try:
            p = doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()
        except KeyError:
            # An empty template may have no list styles — then the paragraph gets a dash.
            p = doc.add_paragraph()
            node.runs.insert(0, ('— ', False, False))

    cfg = _STYLE.get(node.kind, _STYLE['p'])
    pf = p.paragraph_format
    pf.space_before = Pt(cfg['before'])
    pf.space_after = Pt(cfg['after'])

    for text, bold, under in node.runs:
        if not text:
            continue
        if text == _BR:
            p.add_run().add_break()
            continue
        text = strip_emoji(text, trim=False)
        if not text:
            continue
        run = p.add_run(text)
        run.bold = bold or cfg['bold']
        run.underline = under
        if cfg['size']:
            run.font.size = Pt(cfg['size'])
    return p


# Matches the generated "Bundle composition" heading. The Ukrainian form "Склад комплекту" is language data:
# it is what instructions stored before the translation contain.
_COMP_HEAD_RE = re.compile(r'^\s*(?:склад\s+комплекту|bundle\s+composition)\b', re.IGNORECASE)


def _drop_generated_parts(nodes):
    """Strip from the body what we GENERATE ourselves: the title and the composition block.

    AI normalization cuts them out of the text, but until it has been run over all
    instructions, printing must cope with unstripped text too - otherwise the sheet
    would get two titles and two composition lists.

    Title: only the FIRST node, and only if it is an h1: a bold line inside the text
    may be a treatment name, not the sheet title.
    Composition: from the "Bundle composition" heading TO THE END:
    in all instructions this block comes last, and there is no meaningful text after it.
    """
    out = list(nodes)
    if out and out[0].kind == 'h1':
        out = out[1:]
    for i, n in enumerate(out):
        if n.kind in ('h1', 'h2') and _COMP_HEAD_RE.match(n.visible):
            return out[:i]
    return out


def _title_node(sheet):
    """Sheet title = the bundle name from the catalog, as on the site.

    That is exactly how the product is labeled on the storefront, and the buyer
    recognizes what they ordered.
    """
    name = sheet.bundle_name or ''
    node = _Node('h1')
    node.runs.append((strip_emoji(name), True, False))
    return node


def build(plan, excluded_item_ids=(), *, user=None):
    """Build a .docx with instructions. Returns (bytes|None, report).

    An empty result (None) is not an error: the selected orders simply contain no
    bundles.
    """
    from docx import Document
    from docx.enum.section import WD_SECTION

    from catalog import instruction_resolve as ir
    from . import instruction_print

    sheets = instruction_print.collect_sheets(plan, excluded_item_ids, user=user)
    if not sheets:
        return None, {'printed': [], 'placeholders': [], 'substituted': [], 'sheets': 0}

    doc = Document()
    _setup(doc)

    # Instructions that are not ready do NOT go into the file: an empty sheet adds
    # nothing, but wastes paper and print queue time. The list of what has to be
    # inserted by hand stays in the export result panel.
    ready = [sh for sh in sheets if sh.res.kind == ir.KIND_OK]
    placeholders = [sh for sh in sheets if sh.res.kind != ir.KIND_OK]
    if not ready:
        return None, {'printed': [], 'sheets': 0, 'placeholders': _placeholders(placeholders),
                      'substituted': [], 'oversize': [], 'blanks': 0, 'duplex': False,
                      'kits': len(sheets)}

    # We print in REVERSE order: this printer puts the last printed page on top of the
    # stack, so the stack taken from the tray ends up reversed (the last sheet of the
    # file is physically first). With the docx write order reversed, after printing the
    # stack comes out as intended: sheet No. 1 on top, then by order numbers (verified
    # by the owner by hand).
    first = True
    for sh in reversed(ready):
        if not first:
            # SECTION break starting on an odd page, not a plain page break: this way Word
            # itself adds a blank back side after an instruction with an odd number of pages,
            # and each following one lands on the front of a new sheet.
            doc.add_section(WD_SECTION.ODD_PAGE)
        first = False

        # Title and composition are GENERATED, not taken from the text. Both come from the
        # catalog, so they cannot diverge from the name on the site or from the box: if the
        # bundle composition changes, the next print is already correct, with no text edit.
        # The instruction text itself does not contain them (AI normalization cuts them out).
        _write_node(doc, _title_node(sh))

        for node in _drop_generated_parts(parse_html(sh.res.html)):
            _write_node(doc, node)

        try:
            comp = ir.composition_block(sh.res.components)
        except Exception:
            log.exception('composition_block failed for item=%s', sh.item.pk)
            comp = ''
        if comp:
            for node in parse_html(comp):
                _write_node(doc, node)
    printed = ready

    _setup(doc)   # once more: new sections were created after the first call
    buf = BytesIO()
    doc.save(buf)

    report = {
        'printed': printed,
        'sheets': len(ready),
        'placeholders': _placeholders(placeholders),
        'oversize': [],
        'blanks': 0,
        # The file is built with sections that start on an odd page, so with two or more
        # instructions it may be printed ONLY duplex: otherwise the back of one would land
        # on the front of the next. With a single instruction there are no breaks, and the
        # print mode no longer matters.
        'duplex': len(ready) > 1,
        'kits': len(sheets),
        # Sheets where the machine substituted the title (and sometimes the dose). They are
        # printed along with the rest, but the manager must see this: the only place where
        # they learn that the dose could not be recalculated is this line.
        'substituted': _placeholders([sh for sh in ready if sh.res.source == 'auto']),
    }
    return buf.getvalue(), report


def _placeholders(sheets):
    return [
        {'index': sh.index, 'order_number': sh.order.order_number,
         'bundle': sh.bundle_name, 'reason': sh.res.reason,
         'summary': sh.res.summary,
         'override_id': sh.res.override.pk if sh.res.override else None}
        for sh in sheets
    ]


def build_single_bundle(bundle) -> bytes | None:
    """.docx with the instruction of ONE bundle — without an order.

    Unlike build(), it does not go through collect_sheets (that one is tied to
    Order/OrderItem): it resolves the instruction for the bundle's OWN contents, so it
    writes nothing to the DB (no BundleInstructionOverride is created — comparing
    the template with itself never gives discrepancies).
    """
    from docx import Document

    from catalog import instruction_resolve as ir

    res = ir.resolve_bundle(bundle)
    if res.kind != ir.KIND_OK or not res.html.strip():
        return None

    doc = Document()
    _setup(doc)

    title = _Node('h1')
    title.runs.append((strip_emoji(bundle.name), True, False))
    _write_node(doc, title)

    for node in _drop_generated_parts(parse_html(res.html)):
        _write_node(doc, node)

    try:
        comp = ir.composition_block(res.components)
    except Exception:
        log.exception('composition_block failed for bundle=%s', bundle.pk)
        comp = ''
    if comp:
        for node in parse_html(comp):
            _write_node(doc, node)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
